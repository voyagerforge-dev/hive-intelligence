# tests/test_striptrim.py
from pathlib import Path
import docx
from hiveprep.striptrim import strip_struck

def _make_doc(path: Path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("keep this ")
    struck = p.add_run("DELETE THIS")
    struck.font.strike = True
    p.add_run(" and keep this too")
    # struck run inside a table cell
    t = d.add_table(rows=1, cols=1)
    cell_p = t.cell(0, 0).paragraphs[0]
    cell_p.add_run("cell keep ")
    cs = cell_p.add_run("cell DELETE")
    cs.font.strike = True
    d.save(path)

def test_strip_struck_removes_struck_text(tmp_path):
    src = tmp_path / "in.docx"; out = tmp_path / "out.docx"
    _make_doc(src)
    stats = strip_struck(src, out)
    text = "\n".join(p.text for p in docx.Document(out).paragraphs)
    table_text = docx.Document(out).tables[0].cell(0, 0).text
    assert "DELETE THIS" not in text
    assert "keep this" in text and "keep this too" in text
    assert "cell DELETE" not in table_text
    assert "cell keep" in table_text
    assert stats.runs_removed == 2

def test_strip_struck_noop_when_clean(tmp_path):
    src = tmp_path / "c.docx"; out = tmp_path / "co.docx"
    d = docx.Document(); d.add_paragraph().add_run("all good"); d.save(src)
    stats = strip_struck(src, out)
    assert stats.runs_removed == 0
    assert "all good" in docx.Document(out).paragraphs[0].text
