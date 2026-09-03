"""Remove struck-through text from a DOCX, deterministically.

Strikethrough in Word is a run property: <w:strike/> (single) or <w:dstrike/>
(double). We drop any run whose rPr carries either, in body paragraphs AND in
table cells, then save a cleaned copy. This runs BEFORE rendering to PDF so
obsolete content never reaches the vision pass.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import docx
from docx.oxml.ns import qn


@dataclass
class StripStats:
    runs_removed: int = 0


def _run_is_struck(run) -> bool:
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return False
    for tag in ("w:strike", "w:dstrike"):
        el = rpr.find(qn(tag))
        if el is not None and el.get(qn("w:val")) not in ("false", "0"):
            return True
    return False


def _strip_paragraph(paragraph, stats: StripStats) -> None:
    for run in list(paragraph.runs):
        if _run_is_struck(run):
            run._element.getparent().remove(run._element)
            stats.runs_removed += 1


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def strip_struck(src: Path, out: Path) -> StripStats:
    """Load DOCX at `src`, remove all struck runs, save to `out`. Returns stats."""
    document = docx.Document(str(src))
    stats = StripStats()
    for paragraph in document.paragraphs:
        _strip_paragraph(paragraph, stats)
    for table in document.tables:
        for paragraph in _iter_table_paragraphs(table):
            _strip_paragraph(paragraph, stats)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    return stats
